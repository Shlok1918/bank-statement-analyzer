#!/usr/bin/env python3
"""
Bank statement analyzer
=======================
Reads a (password-protected) bank-statement .xlsx, cleans it, categorizes every
transaction, works out who you paid / who paid you, and produces:

  1. <outdir>/bank_report.xlsx   - multi-sheet Excel report (filterable tables)
  2. <outdir>/dashboard.html     - interactive dashboard (open in any browser)
  3. <outdir>/transactions_clean.csv

Install once:
    pip install pandas openpyxl msoffcrypto-tool

Run:
    python bank_analyzer.py statement.xlsx
    python bank_analyzer.py statement.xlsx -p "yourPassword" -o my_report

(If you don't pass -p it asks for the password without echoing it, or reads the
BANK_XLSX_PASSWORD environment variable.)
"""
import argparse
import getpass
import io
import json
import os
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import msoffcrypto
except ImportError:
    sys.exit("Missing dependency. Run:  pip install pandas openpyxl msoffcrypto-tool")

# =============================================================================
# USER CONFIG  - edit these to improve accuracy for YOUR statement
# =============================================================================
# Your own names / other own-account names. Transfers whose narration contains
# one of these are tagged "Self Transfer" (and can be excluded from spending).
MY_NAMES = []  # e.g. ["RAHUL SHARMA", "R SHARMA"]

# Your rules run FIRST and override everything else.
# (category, regex matched against the narration, side)   side: debit | credit | any
CUSTOM_RULES = [
    # ("Rent & Housing", r"MR LANDLORD NAME", "debit"),
    # ("Salary",         r"YOUR EMPLOYER NAME", "credit"),
]

# Exclude "Self Transfer" and "Credit Card Payment" from spend/income totals so
# money moving between your own accounts is not counted twice.
EXCLUDE_FROM_TOTALS = {"Self Transfer", "Credit Card Payment"}

# =============================================================================
# CATEGORY RULES  (first match wins, order matters)
# =============================================================================
DEBIT_RULES = [
    ("Bank Charges & Taxes", r"CHARGES|CHRG|\bFEE\b|\bFEES\b.*(CARD|BANK)|\bGST\b|SMS ALERT|ANNUAL FEE|PENALTY|MIN(IMUM)? BAL|\bAMC\b|\bTDS\b|\bCESS\b|STAMP DUTY"),
    ("Credit Card Payment", r"CREDIT CARD|\bCC PAYMENT|CRED CLUB|CREDCLUB|\bCRED\b|CARD PAYMENT|CARDS? BILL|BILLDESK.*CARD|AUTOPAY.*CARD"),
    ("Cash Withdrawal", r"\bATM\b|\bNWD\b|CASH WDL|CASH WITHDRAWAL|\bCWDR\b|\bATW\b|ATM-CASH"),
    ("Investments", r"ZERODHA|GROWW|UPSTOX|ANGEL ?(ONE|BROKING)|KUVERA|\bCOIN\b|MUTUAL FUND|\bSIP\b|\bMF\b|NSDL|CDSL|INDMONEY|SMALLCASE|\bPPF\b|\bNPS\b|NSE CLEARING|ICCL|BSE LTD|STOCK|FD BOOKING|\bRD INSTALL"),
    ("Insurance", r"INSURANCE|\bLIC\b|POLICYBAZAAR|\bACKO\b|\bDIGIT\b|STAR HEALTH|HDFC LIFE|ICICI PRU|PREMIUM|MAX LIFE|BAJAJ ALLIANZ"),
    ("EMI & Loans", r"\bEMI\b|\bLOAN\b|BAJAJ FIN|HDFC LTD|KREDITBEE|MONEYVIEW|ZESTMONEY|LAZYPAY|\bSIMPL\b|\bSLICE\b|FINANCE|FINSERV|HOME FIN|CAPITAL FIRST|MPOKKET|CASHE"),
    ("Rent & Housing", r"\bRENT\b|LANDLORD|HOUSING|MAINTENANCE|SOCIETY|NOBROKER|BROKERAGE|NESTAWAY"),
    ("Food Delivery & Dining", r"SWIGGY|ZOMATO|DOMINO|PIZZA|MCDONALD|\bKFC\b|BURGER|STARBUCKS|CAFE|RESTAURANT|BARISTA|BEHROUZ|FAASOS|EATSURE|EATFIT|SUBWAY|CHAAYOS|BAKERY|DHABA|BIRYANI|HOTEL|SWEETS|FOOD|KITCHEN|TEA|COFFEE|JUICE|ICE ?CREAM|HALDIRAM|WAGHMARE|MAGICPIN"),
    ("Groceries", r"BIGBASKET|BLINKIT|GROFERS|ZEPTO|INSTAMART|DMART|AVENUE SUPERMARTS|RELIANCE (FRESH|SMART|RETAIL)|JIOMART|MILK|KIRANA|SUPERMARKET|SUPER MARKET|FRUIT|VEGETABLE|MORE RETAIL|NATURES BASKET|DAIRY|PROVISION|GENERAL STORE|MART\b|DUNZO"),
    ("Shopping", r"AMAZON|FLIPKART|MYNTRA|AJIO|MEESHO|NYKAA|TATA ?CLIQ|DECATHLON|LENSKART|\bIKEA\b|CROMA|RELIANCE DIGITAL|SNAPDEAL|FIRSTCRY|SHOPPERS STOP|WESTSIDE|PANTALOONS|LIFESTYLE|MAX FASHION|FASHION|CLOTH|FOOTWEAR|ELECTRONICS|MOBILE|STORES?\b|ENTERPRISE|TRADERS"),
    ("Fuel", r"PETROL|DIESEL|HPCL|BPCL|IOCL|INDIAN ?OIL|\bSHELL\b|NAYARA|RELIANCE BP|FUEL|\bCNG\b|FILLING STATION"),
    ("Travel & Transport", r"\bUBER\b|\bOLA\b|OLACABS|RAPIDO|IRCTC|REDBUS|MAKEMYTRIP|\bMMT\b|GOIBIBO|IXIGO|CLEARTRIP|INDIGO|AIR ?INDIA|VISTARA|AKASA|SPICEJET|METRO|FASTAG|NHAI|BEST UNDERTAKING|ABHIBUS|YULU|BLUSMART|\bRTO\b|PARKING|TOLL|TRAVELS?\b|CAB\b|TAXI|MUMBAI ONE|MMRDA"),
    ("Telecom & Recharge", r"\bJIO\b|AIRTEL|VODAFONE|\bVI\b|\bBSNL\b|RECHARGE|PREPAID|\bDTH\b|TATA ?PLAY|TATASKY|DISH ?TV|ACT FIBER|HATHWAY|BROADBAND"),
    ("Bills & Utilities", r"ELECTRIC|MSEDCL|MSEB|ADANI ELEC|TATA POWER|\bBSES\b|BESCOM|WATER BILL|\bGAS\b|MAHANAGAR|\bMGL\b|PIPED|BBPS|BILLDESK|BILL PAY|PAYU|UTILITY|MUNICIPAL|PROPERTY TAX"),
    ("Entertainment & Subscriptions", r"NETFLIX|SPOTIFY|HOTSTAR|DISNEY|PRIME ?VIDEO|YOUTUBE|GOOGLE ?(PLAY|ONE|CLOUD)|APPLE|ICLOUD|BOOKMYSHOW|\bPVR\b|\bINOX\b|SONYLIV|ZEE5|JIOCINEMA|STEAM|CHATGPT|OPENAI|ANTHROPIC|CLAUDE|GITHUB|NOTION|CANVA|ADOBE|MICROSOFT|LINKEDIN|DROPBOX|AUDIBLE|KINDLE|DREAM11|MPL|GAMING|PLAYSTORE"),
    ("Health & Fitness", r"PHARMA|MEDICAL|MEDICINE|MEDPLUS|APOLLO|\b1MG\b|NETMEDS|PHARMEASY|HOSPITAL|CLINIC|DIAGNOSTIC|\bLABS?\b|PRACTO|CULT\.?FIT|\bGYM\b|DOCTOR|DENTAL|THYROCARE|LAL ?PATH|HEALTH|CHEMIST|SALON|SPA\b"),
    ("Education", r"UDEMY|COURSERA|UNACADEMY|BYJU|SCHOOL|COLLEGE|TUITION|EDUCATION|UNIVERSITY|EXAM|UPGRAD|SIMPLILEARN|INSTITUTE|ACADEMY|\bFEES?\b"),
    ("Donations & Gifts", r"DONATION|TEMPLE|\bNGO\b|CHARITY|MANDIR|PM ?CARES|TRUST\b"),
]

CREDIT_RULES = [
    ("Salary", r"SALARY|\bSAL\b|PAYROLL|STIPEND|WAGES"),
    ("Interest Earned", r"INTEREST|\bINT\.? ?PD\b|\bINT\.? ?CREDIT|INT ON|SB INT"),
    ("Refund & Cashback", r"REFUND|REVERSAL|CASHBACK|\bREV\b|\bREVERS|RETURN|FAILED|CHARGEBACK"),
    ("Investment Returns", r"DIVIDEND|REDEMPTION|MATURITY|ZERODHA|GROWW|UPSTOX|MUTUAL FUND|\bMF\b|NSE CLEARING|ICCL|\bFD\b"),
    ("Cash Deposit", r"CASH DEP|\bCDM\b|CASH DEPOSIT|BY CASH"),
]

SELF_PATTERNS = r"\bSELF\b|OWN ACC|OWN A/C|SELF TRANSFER|TRF TO SELF|TO SELF"

# Canonical display names so "SWIGGY BANGALORE", "Swiggy Instamart" etc. group
BRANDS = {
    r"SWIGGY": "Swiggy", r"ZOMATO": "Zomato", r"AMAZON|AMZN": "Amazon", r"FLIPKART": "Flipkart",
    r"MYNTRA": "Myntra", r"AJIO": "Ajio", r"MEESHO": "Meesho", r"NYKAA": "Nykaa",
    r"BLINKIT|GROFERS": "Blinkit", r"ZEPTO": "Zepto", r"BIGBASKET": "BigBasket", r"DMART|AVENUE SUPERMARTS": "DMart",
    r"\bUBER\b": "Uber", r"\bOLA\b|OLACABS": "Ola", r"RAPIDO": "Rapido", r"IRCTC": "IRCTC", r"REDBUS": "RedBus",
    r"MAKEMYTRIP": "MakeMyTrip", r"NETFLIX": "Netflix", r"SPOTIFY": "Spotify", r"HOTSTAR|DISNEY": "Disney+ Hotstar",
    r"YOUTUBE": "YouTube", r"GOOGLE ?PLAY|PLAYSTORE": "Google Play", r"APPLE|ICLOUD": "Apple", r"BOOKMYSHOW": "BookMyShow",
    r"\bJIO\b": "Jio", r"AIRTEL": "Airtel", r"VODAFONE|\bVI\b": "Vi", r"ZERODHA": "Zerodha", r"GROWW": "Groww",
    r"PHONEPE": "PhonePe", r"PAYTM": "Paytm", r"CRED CLUB|CREDCLUB|\bCRED\b": "CRED", r"STARBUCKS": "Starbucks",
    r"DOMINO": "Domino's", r"MCDONALD": "McDonald's", r"\bKFC\b": "KFC", r"DECATHLON": "Decathlon",
    r"CHATGPT|OPENAI": "OpenAI", r"ANTHROPIC|CLAUDE": "Anthropic", r"GITHUB": "GitHub", r"MICROSOFT": "Microsoft",
    r"APOLLO": "Apollo", r"PHARMEASY": "PharmEasy", r"\b1MG\b": "Tata 1mg", r"LENSKART": "Lenskart",
}

# Heuristics for UPI "merchant vs person"
MERCHANT_HINT = re.compile(
    r"PAYTMQR|PAYTM\.S|BHARATPE|RAZORPAY|PAYU|CASHFREE|BILLDESK|MERCHANT|\bQR\b|GPAY-\d+|^\d{6,}@|^Q\d+@|"
    r"STORE|MART|ENTERPRISE|TRADERS|SERVICES|FOODS|SWEETS|KIRANA|AGENC|PVT|LTD|LLP|CAFE|HOTEL|RESTAURANT|"
    r"PHARMA|MEDICAL|CLINIC|SALON|BAKERY|DAIRY|WHOLESALE|GENERAL|STUDIO|FASHION|WEAR|TAILOR|PETROL|GARAGE|"
    r"MOTORS|ELECTRONICS|MOBILE|BOOK|STATIONERY|SHOP|BHANDAR|VASTRA|JEWEL|RETAIL|SUPER|FRUIT|VEG|CHEMIST|"
    r"TEA|COFFEE|JUICE|PIZZA|BURGER|FOOD|KITCHEN|RESTO|DHABA|CANTEEN|MESS|GYM|LAUNDRY|DRY ?CLEAN|PARLOUR|PARLOR",
    re.I,
)

STOP = {
    "UPI", "IMPS", "NEFT", "RTGS", "DR", "CR", "TRANSFER", "TO", "FROM", "BY", "PAYMENT", "PAYMENTS", "PAY", "TXN",
    "REF", "NO", "DEBIT", "CREDIT", "PAID", "UPIINTENT", "UPIQR", "MOB", "MOBILE", "BANKING", "INB", "NETBANKING",
    "P2A", "P2M", "P2P", "UTR", "RRN", "TPT", "MMT", "FT", "OK", "SENT", "RECEIVED", "COLLECT", "AUTOPAY", "MANDATE",
    "POS", "ECOM", "ATM", "NWD", "ACH", "NACH", "CHQ", "CHEQUE", "CLG", "INF", "INFT", "BIL", "ONL", "ONLINE", "THE",
    "AND", "FOR", "DEP", "WDL", "WITHDRAWAL", "DEPOSIT", "SELF", "MR", "MRS", "MS", "SHRI", "SMT",
}
BANK_CODES = {
    "HDFC", "ICICI", "SBI", "AXIS", "KOTAK", "YES", "YESB", "PNB", "BOB", "CANARA", "IDFC", "INDUSIND", "FEDERAL",
    "UNION", "BANDHAN", "PAYTM", "IPPB", "AIRTEL", "BANK", "OKAXIS", "OKHDFCBANK", "OKICICI", "OKSBI", "YBL", "IBL",
    "AXL", "APL", "PTYES", "PTSBI", "PTAXIS", "HDFCBANK", "ICICIBANK", "KMBL", "UBIN", "SBIN", "BARB", "CNRB",
}
IFSC = re.compile(r"^[A-Z]{4}0[A-Z0-9]{6}$")
VPA = re.compile(r"[\w.\-]+@[a-zA-Z]+")


# =============================================================================
# 1. LOAD (decrypt + find header + normalise)
# =============================================================================
def open_workbook_bytes(path: Path, password: str) -> io.BytesIO:
    out = io.BytesIO()
    with open(path, "rb") as f:
        try:
            office = msoffcrypto.OfficeFile(f)
        except Exception:
            # not an OLE/encrypted container -> plain xlsx
            f.seek(0)
            return io.BytesIO(f.read())
        if office.is_encrypted():
            if not password:
                sys.exit("This file is password protected - please provide the password.")
            try:
                office.load_key(password=password)
                office.decrypt(out)
            except Exception as e:
                sys.exit(f"Could not decrypt the file (wrong password?): {e}")
        else:
            f.seek(0)
            out.write(f.read())
    out.seek(0)
    return out


def pick(cols, patterns, used):
    for pat in patterns:
        for c in cols:
            if c not in used and re.search(pat, c.lower()):
                used.add(c)
                return c
    return None


def find_header_row(raw: pd.DataFrame):
    for i in range(min(80, len(raw))):
        txt = " | ".join(str(x).lower() for x in raw.iloc[i].tolist() if pd.notna(x))
        if "date" in txt and re.search(r"narration|description|particulars|details|remarks", txt) \
                and re.search(r"debit|withdraw|credit|deposit|amount|\bdr\b|\bcr\b", txt):
            return i
    return None


def to_num(s: pd.Series) -> pd.Series:
    s = s.astype(str).str.replace(r"[^\d.\-]", "", regex=True)
    return pd.to_numeric(s, errors="coerce")


def parse_dates(s: pd.Series) -> pd.Series:
    if pd.api.types.is_datetime64_any_dtype(s):
        return s
    num = pd.to_numeric(s, errors="coerce")
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    mask_serial = num.between(30000, 60000)
    if mask_serial.any():
        out[mask_serial] = pd.to_datetime(num[mask_serial], unit="D", origin="1899-12-30")
    rest = ~mask_serial
    try:
        out[rest] = pd.to_datetime(s[rest].astype(str).str.strip(), dayfirst=True, errors="coerce", format="mixed")
    except (TypeError, ValueError):
        out[rest] = pd.to_datetime(s[rest].astype(str).str.strip(), dayfirst=True, errors="coerce")
    return out


def load_statement(buf: io.BytesIO, only_sheet=None) -> pd.DataFrame:
    xl = pd.ExcelFile(buf)
    frames = []
    for sheet in xl.sheet_names:
        if only_sheet and sheet != only_sheet:
            continue
        raw = xl.parse(sheet, header=None, dtype=object)
        hi = find_header_row(raw)
        if hi is None:
            print(f"  - skipping sheet '{sheet}' (no transaction header found)")
            continue
        header = [str(x).strip() if pd.notna(x) else f"col{j}" for j, x in enumerate(raw.iloc[hi])]
        df = raw.iloc[hi + 1:].copy()
        df.columns = header
        cols = list(df.columns)
        used = set()
        c_drcr = pick(cols, [r"dr\s*/\s*cr|cr\s*/\s*dr|debit\s*/\s*credit|credit\s*/\s*debit|^type$|txn type"], used)
        c_date = pick(cols, [r"(txn|transaction|tran|posting)\s*date", r"^date", r"^(?!.*value).*date", r"date"], used)
        c_desc = pick(cols, [r"narration", r"description", r"particulars", r"details", r"remarks"], used)
        c_ref = pick(cols, [r"chq|cheque|ref"], used)
        c_deb = pick(cols, [r"withdraw", r"debit", r"^dr\b", r"\bdr\b"], used)
        c_cre = pick(cols, [r"deposit", r"credit", r"^cr\b", r"\bcr\b"], used)
        c_amt = pick(cols, [r"^amount", r"amount"], used) if not (c_deb and c_cre) else None
        c_bal = pick(cols, [r"balance"], used)
        if not (c_date and c_desc and ((c_deb and c_cre) or c_amt)):
            print(f"  - skipping sheet '{sheet}' (could not map columns: {cols})")
            continue

        out = pd.DataFrame(index=df.index)
        out["date"] = parse_dates(df[c_date])
        out["desc"] = df[c_desc].fillna("").astype(str).str.replace(r"\s+", " ", regex=True).str.strip()
        out["ref"] = df[c_ref].astype(str).str.strip() if c_ref else ""
        if c_deb and c_cre:
            out["debit"] = to_num(df[c_deb])
            out["credit"] = to_num(df[c_cre])
        else:
            amt = to_num(df[c_amt]).abs()
            if c_drcr:
                is_dr = df[c_drcr].astype(str).str.contains(r"^\s*(dr|d|debit|w)", case=False, regex=True)
            else:
                raw_amt = df[c_amt].astype(str)
                is_dr = raw_amt.str.contains(r"\bDr\b|^\s*-|\(", case=False, regex=True)
            out["debit"] = amt.where(is_dr)
            out["credit"] = amt.where(~is_dr)
        out["balance"] = to_num(df[c_bal]) if c_bal else np.nan

        # --- stitch multi-line narrations (continuation rows have no date/amount)
        has_amt = out["debit"].notna() | out["credit"].notna()
        valid = out["date"].notna() & has_amt
        cont = ~out["date"].notna() & ~has_amt & out["desc"].ne("") & out["desc"].ne("nan")
        grp = valid.cumsum()
        extra = out.loc[cont & (grp > 0), "desc"].groupby(grp[cont & (grp > 0)]).agg(" ".join)
        out = out[valid].copy()
        g = grp[valid]
        out["desc"] = (out["desc"] + " " + g.map(extra).fillna("")).str.strip().str.slice(0, 400)
        out["sheet"] = sheet
        frames.append(out)
        print(f"  + sheet '{sheet}': {len(out)} transactions")

    if not frames:
        sys.exit("No transaction table found. Check the file / pass --sheet.")
    df = pd.concat(frames, ignore_index=True)
    df["debit"] = df["debit"].fillna(0.0)
    df["credit"] = df["credit"].fillna(0.0)
    df = df[(df["debit"] > 0) | (df["credit"] > 0)]
    if df["balance"].notna().any():  # overlapping exports -> drop exact duplicates
        before = len(df)
        df = df.drop_duplicates(subset=["date", "desc", "debit", "credit", "balance"])
        if before != len(df):
            print(f"  - removed {before - len(df)} duplicate rows (overlapping sheets)")
    df = df.sort_values("date", kind="stable").reset_index(drop=True)
    return df


# =============================================================================
# 2. ENRICH (channel, counterparty, category)
# =============================================================================
def detect_channel(u: str) -> str:
    for name, pat in [
        ("UPI", r"\bUPI\b"), ("IMPS", r"\bIMPS\b|\bMMT\b"), ("NEFT", r"\bNEFT\b"), ("RTGS", r"\bRTGS\b"),
        ("ATM / Cash", r"\bATM\b|\bNWD\b|CASH WDL|CASH WITHDRAWAL|CASH DEP|\bCDM\b"),
        ("Card / POS", r"\bPOS\b|\bECOM\b|DEBIT CARD|CREDIT CARD|\bVISA\b|MASTERCARD|RUPAY"),
        ("Auto-debit (ACH/NACH)", r"\bACH\b|\bNACH\b|\bECS\b|MANDATE|AUTOPAY"),
        ("Cheque", r"\bCHQ\b|CHEQUE|\bCLG\b|\bINF\b"),
        ("Interest / Charges", r"INTEREST|CHARGES|\bINT\.?PD|\bFEE\b|\bGST\b"),
        ("Net Banking", r"NET ?BANKING|\bINB\b|\bIB\b|BILLDESK|\bBBPS\b"),
    ]:
        if re.search(pat, u):
            return name
    return "Other"


def extract_party(narr: str, channel: str):
    """Return (party name, vpa)"""
    u = narr.upper()
    m = VPA.search(narr)
    vpa = m.group(0).lower() if m else ""

    if channel == "ATM / Cash":
        return "ATM / Cash", vpa
    if channel == "Interest / Charges":
        return "Bank (Interest / Charges)", vpa

    # match brands on the narration WITHOUT the UPI handle (so "@paytm"/"@ybl" don't hijack the name),
    # then on the part of the UPI id before the "@"
    u_no_vpa = VPA.sub(" ", u)
    for pat, brand in BRANDS.items():
        if re.search(pat, u_no_vpa):
            return brand, vpa
        if vpa and brand not in ("Paytm", "PhonePe") and re.search(pat, vpa.split("@")[0].upper()):
            return brand, vpa

    if channel in ("UPI", "IMPS", "NEFT", "RTGS", "Net Banking", "Cheque"):
        tokens = [t.strip() for t in re.split(r"[/\-]", narr) if t.strip()]
        for t in tokens:
            tu = t.upper()
            if "@" in t or IFSC.match(tu.replace(" ", "")):
                continue
            letters = sum(ch.isalpha() or ch in " .&'" for ch in t)
            if len(t) < 3 or letters / max(len(t), 1) < 0.75:
                continue
            words = [w for w in re.split(r"\s+", tu) if w]
            words = [w for w in words if w not in STOP]
            if not words:
                continue
            name = " ".join(words)
            if name in BANK_CODES or name.endswith(" BANK") or name.replace(" ", "") in BANK_CODES:
                continue
            return re.sub(r"\s+", " ", name).strip().title(), vpa
        if vpa:
            return vpa.split("@")[0].replace(".", " ").replace("_", " ").title(), vpa

    # card / POS / others: first 1-2 sensible words
    words = [w for w in re.split(r"[\s/\-]+", u)
             if len(w) >= 3 and w not in STOP and sum(c.isalpha() for c in w) / len(w) >= 0.85
             and w not in BANK_CODES]
    if words:
        return " ".join(words[:2]).title(), vpa
    return "Unknown", vpa


def categorize(narr: str, party: str, vpa: str, channel: str, is_debit: bool) -> str:
    u = narr.upper()
    side = "debit" if is_debit else "credit"
    for cat, pat, s in CUSTOM_RULES:
        if s in ("any", side) and re.search(pat, u, re.I):
            return cat
    if MY_NAMES and any(n.upper() in u for n in MY_NAMES):
        return "Self Transfer"
    if re.search(SELF_PATTERNS, u):
        return "Self Transfer"

    rules = DEBIT_RULES if is_debit else CREDIT_RULES
    for cat, pat in rules:
        if re.search(pat, u):
            return cat
    # party name might reveal the category (e.g. "Sharma General Store")
    for cat, pat in rules:
        if re.search(pat, party.upper()):
            return cat

    if channel == "UPI":
        is_merchant = bool(MERCHANT_HINT.search(party + " " + vpa))
        if is_debit:
            return "UPI - Merchant / Shop Payments" if is_merchant else "UPI - Person to Person (Sent)"
        return "UPI - Business Receipts" if is_merchant else "UPI - Person to Person (Received)"
    if channel in ("IMPS", "NEFT", "RTGS"):
        return "Bank Transfer Out" if is_debit else "Bank Transfer In"
    if channel == "Card / POS":
        return "Card / POS - Other"
    if channel == "Cheque":
        return "Cheque Payments" if is_debit else "Cheque Deposits"
    return "Other Debits" if is_debit else "Other Credits"


def enrich(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df["channel"] = df["desc"].str.upper().map(detect_channel)
    parts = [extract_party(d, c) for d, c in zip(df["desc"], df["channel"])]
    df["party"] = [p[0] for p in parts]
    df["vpa"] = [p[1] for p in parts]
    df["type"] = np.where(df["debit"] > 0, "Debit", "Credit")
    df["amount"] = np.where(df["debit"] > 0, df["debit"], df["credit"])
    df["category"] = [
        categorize(d, p, v, c, t == "Debit")
        for d, p, v, c, t in zip(df["desc"], df["party"], df["vpa"], df["channel"], df["type"])
    ]
    df["month"] = df["date"].dt.to_period("M").astype(str)
    df["year"] = df["date"].dt.year
    df["weekday"] = df["date"].dt.day_name()
    df["counted"] = ~df["category"].isin(EXCLUDE_FROM_TOTALS)
    return df


# =============================================================================
# 3. AGGREGATES
# =============================================================================
def build_tables(df: pd.DataFrame) -> dict:
    t = {}
    c = df[df["counted"]]

    m = c.groupby("month").agg(Received=("credit", "sum"), Spent=("debit", "sum"), Txns=("desc", "count"))
    m["Net"] = m["Received"] - m["Spent"]
    m["Savings Rate"] = np.where(m["Received"] > 0, m["Net"] / m["Received"], np.nan)
    if df["balance"].notna().any():
        m["Closing Balance"] = df.dropna(subset=["balance"]).groupby("month")["balance"].last()
    t["Monthly"] = m.reset_index().rename(columns={"month": "Month"})

    sp = c[c["debit"] > 0].groupby("category").agg(Total=("debit", "sum"), Txns=("debit", "count"), Avg=("debit", "mean"),
                                                      Largest=("debit", "max"))
    sp["% of Spend"] = sp["Total"] / sp["Total"].sum()
    t["Spend by Category"] = sp.sort_values("Total", ascending=False).reset_index().rename(columns={"category": "Category"})

    inc = c[c["credit"] > 0].groupby("category").agg(Total=("credit", "sum"), Txns=("credit", "count"), Avg=("credit", "mean"),
                                                       Largest=("credit", "max"))
    inc["% of Income"] = inc["Total"] / inc["Total"].sum()
    t["Income by Category"] = inc.sort_values("Total", ascending=False).reset_index().rename(columns={"category": "Category"})

    piv = c[c["debit"] > 0].pivot_table(index="category", columns="month", values="debit", aggfunc="sum", fill_value=0)
    piv["TOTAL"] = piv.sum(axis=1)
    t["Category x Month (Spend)"] = piv.sort_values("TOTAL", ascending=False).reset_index().rename(columns={"category": "Category"})

    def party_table(col, side_df):
        g = side_df.groupby("party").agg(
            Total=(col, "sum"), Txns=(col, "count"), Avg=(col, "mean"), Largest=(col, "max"),
            First=("date", "min"), Last=("date", "max"),
            Main_Category=("category", lambda s: s.mode().iat[0]),
            Months_Active=("month", "nunique"),
        ).sort_values("Total", ascending=False).reset_index().rename(columns={"party": "Party", "Main_Category": "Main Category",
                                                                              "Months_Active": "Months Active"})
        g["First"] = g["First"].dt.date
        g["Last"] = g["Last"].dt.date
        return g

    t["Paid To (Payees)"] = party_table("debit", c[c["debit"] > 0])
    t["Received From (Payers)"] = party_table("credit", c[c["credit"] > 0])

    ch = c.groupby("channel").agg(Spent=("debit", "sum"), Received=("credit", "sum"), Txns=("desc", "count"))
    t["By Channel"] = ch.sort_values("Spent", ascending=False).reset_index().rename(columns={"channel": "Channel"})

    # recurring payments: same payee, >=3 distinct months, fairly stable amount
    rec = []
    for p, g in c[c["debit"] > 0].groupby("party"):
        months = g["month"].nunique()
        if months >= 3 and p != "Unknown":
            mean, std = g["debit"].mean(), g["debit"].std(ddof=0)
            if mean > 0 and std / mean <= 0.15:
                rec.append({"Payee": p, "Category": g["category"].mode().iat[0], "Typical Amount": round(mean, 2),
                            "Months Seen": months, "Total Paid": g["debit"].sum(), "Last Paid": g["date"].max().date()})
    t["Recurring Payments"] = pd.DataFrame(rec).sort_values("Total Paid", ascending=False) if rec else pd.DataFrame(
        columns=["Payee", "Category", "Typical Amount", "Months Seen", "Total Paid", "Last Paid"])

    # things the rules couldn't place -> use this to extend CUSTOM_RULES
    rev = df[df["category"].isin(["Other Debits", "Other Credits", "Card / POS - Other", "Bank Transfer Out", "Bank Transfer In",
                                   "UPI - Person to Person (Sent)", "UPI - Person to Person (Received)"])]
    rv = rev.groupby(["party", "category"]).agg(Total=("amount", "sum"), Txns=("amount", "count"),
                                                Example=("desc", "first")).reset_index().sort_values("Total", ascending=False)
    t["Review & Refine"] = rv.rename(columns={"party": "Party", "category": "Current Category"})
    return t


# =============================================================================
# 4. EXCEL REPORT
# =============================================================================
def write_excel(df: pd.DataFrame, tables: dict, path: Path):
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    c = df[df["counted"]]
    summary = pd.DataFrame({
        "Metric": ["Statement period", "Months covered", "Total received (excl. self/card transfers)",
                   "Total spent (excl. self/card transfers)", "Net saved", "Savings rate", "Avg monthly spend",
                   "Avg monthly income", "Transactions", "Largest single expense", "Largest single credit",
                   "Moved between own accounts / card bills (excluded)"],
        "Value": [f"{df['date'].min():%d %b %Y}  to  {df['date'].max():%d %b %Y}", df["month"].nunique(),
                  c["credit"].sum(), c["debit"].sum(), c["credit"].sum() - c["debit"].sum(),
                  (c["credit"].sum() - c["debit"].sum()) / c["credit"].sum() if c["credit"].sum() else np.nan,
                  c["debit"].sum() / max(df["month"].nunique(), 1), c["credit"].sum() / max(df["month"].nunique(), 1),
                  len(df), c["debit"].max(), c["credit"].max(), df.loc[~df["counted"], "amount"].sum()],
    })

    tx = df[["date", "month", "desc", "party", "category", "channel", "type", "debit", "credit", "balance", "ref", "vpa"]].copy()
    tx["date"] = tx["date"].dt.date
    tx.columns = ["Date", "Month", "Description", "Party", "Category", "Channel", "Type", "Debit", "Credit", "Balance",
                  "Ref", "UPI ID"]

    sheets = {"Summary": summary, "Transactions": tx, **tables}
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, d in sheets.items():
            d.to_excel(xw, sheet_name=name[:31], index=False)

        head_fill = PatternFill("solid", fgColor="1F3A5F")
        for name in sheets:
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            if name != "Summary":
                ws.auto_filter.ref = ws.dimensions
            headers = [cell.value for cell in ws[1]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF", name="Arial")
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for j, h in enumerate(headers, start=1):
                hl = str(h).lower()
                is_count = any(k in hl for k in ("txns", "months", "count"))
                is_pct = "%" in hl or "rate" in hl
                width = max([len(str(h))] + [len(str(x.value)) for x in ws[get_column_letter(j)][1:300] if x.value is not None])
                ws.column_dimensions[get_column_letter(j)].width = min(max(width + 2, 10), 60)
                for cell in ws[get_column_letter(j)][1:]:
                    cell.font = Font(name="Arial", size=10)
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        if is_pct:
                            cell.number_format = "0.0%"
                        elif is_count or hl in ("month", "year"):
                            cell.number_format = "0"
                        else:
                            cell.number_format = "#,##0.00"
                    elif hl == "date" or hl in ("first", "last", "last paid"):
                        cell.number_format = "dd-mmm-yyyy"

        # Summary sheet: format values per-row
        ws = xw.sheets["Summary"]
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 34
        for r in range(2, ws.max_row + 1):
            label = str(ws.cell(r, 1).value)
            cell = ws.cell(r, 2)
            cell.alignment = Alignment(horizontal="right")
            if label == "Savings rate":
                cell.number_format = "0.0%"
            elif label in ("Months covered", "Transactions"):
                cell.number_format = "0"
            elif isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"

        # Native chart on Monthly sheet
        wm = xw.sheets["Monthly"]
        n = wm.max_row
        bar = BarChart()
        bar.title = "Received vs Spent by month"
        bar.add_data(Reference(wm, min_col=2, max_col=3, min_row=1, max_row=n), titles_from_data=True)
        bar.set_categories(Reference(wm, min_col=1, min_row=2, max_row=n))
        bar.height, bar.width = 9, 22
        wm.add_chart(bar, f"A{n + 3}")


# =============================================================================
# 5. HTML DASHBOARD
# =============================================================================
HTML = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Bank Statement Dashboard</title>
<script>window.addEventListener("error",function(e){var b=document.getElementById("err");if(b){b.style.display="block";b.textContent="Dashboard error: "+e.message+" (line "+e.lineno+")"}});</script>
__CHARTJS__
<style>
:root{--bg:#f4f6f9;--card:#fff;--ink:#1b2432;--mute:#677386;--line:#e4e8ef;--in:#16a34a;--out:#dc2626;--acc:#2563eb}
*{box-sizing:border-box}
body{margin:0;font-family:system-ui,-apple-system,"Segoe UI",Roboto,sans-serif;background:var(--bg);color:var(--ink)}
header{display:flex;flex-wrap:wrap;gap:12px;align-items:center;justify-content:space-between;padding:18px 24px}
h1{font-size:20px;margin:0}h2{font-size:14px;margin:0 0 10px;color:var(--mute);font-weight:600;text-transform:uppercase;letter-spacing:.04em}
.sub{color:var(--mute);font-size:13px}
.filters{display:flex;gap:8px;align-items:center;flex-wrap:wrap;font-size:13px}
select,input{padding:7px 9px;border:1px solid var(--line);border-radius:8px;background:#fff;font:inherit}
main{padding:0 24px 40px;max-width:1500px;margin:auto}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px;margin-bottom:14px}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:14px 16px}
.kpi .v{font-size:22px;font-weight:700;margin-top:4px}.kpi .l{font-size:12px;color:var(--mute)}
.in{color:var(--in)}.out{color:var(--out)}
.grid{display:grid;gap:14px;grid-template-columns:repeat(12,1fr);margin-bottom:14px}
.c12{grid-column:span 12}.c8{grid-column:span 8}.c6{grid-column:span 6}.c4{grid-column:span 4}
@media(max-width:1000px){.c8,.c6,.c4{grid-column:span 12}}
.chartbox{position:relative;height:300px}.tall{height:420px}
table{width:100%;border-collapse:collapse;font-size:13px}
th{position:sticky;top:0;background:#f0f3f8;text-align:left;padding:8px;border-bottom:1px solid var(--line);cursor:pointer}
td{padding:7px 8px;border-bottom:1px solid var(--line);vertical-align:top}
td.n,th.n{text-align:right;white-space:nowrap}
.wrap{max-height:520px;overflow:auto;border:1px solid var(--line);border-radius:8px}
.pill{display:inline-block;padding:2px 8px;border-radius:999px;background:#eef2f9;font-size:12px}
button{padding:8px 14px;border:1px solid var(--line);background:#fff;border-radius:8px;cursor:pointer;font:inherit}
</style></head><body>
<header>
  <div><h1>Bank Statement Dashboard</h1><div class="sub" id="period"></div></div>
  <div class="filters">
    <label>From <select id="from"></select></label>
    <label>To <select id="to"></select></label>
    <label><input type="checkbox" id="incself"> include self / card-bill transfers</label>
  </div>
</header>
<div id="err" style="display:none;background:#fee2e2;color:#991b1b;padding:10px 24px;font-size:13px"></div>
<main>
  <div class="kpis" id="kpis"></div>
  <div class="grid">
    <div class="card c8"><h2>Money in vs out, by month</h2><div class="chartbox"><canvas id="ch_month"></canvas></div></div>
    <div class="card c4"><h2>Where the money went</h2><div class="chartbox"><canvas id="ch_cat"></canvas></div></div>
    <div class="card c12"><h2>Spending by category, month by month</h2><div class="chartbox tall"><canvas id="ch_catmonth"></canvas></div></div>
    <div class="card c6"><h2>Top payees (spent)</h2><div class="chartbox tall"><canvas id="ch_payees"></canvas></div></div>
    <div class="card c6"><h2>Top payers (received)</h2><div class="chartbox tall"><canvas id="ch_payers"></canvas></div></div>
    <div class="card c6"><h2>Income by category</h2><div class="chartbox"><canvas id="ch_inc"></canvas></div></div>
    <div class="card c6"><h2>Channel split (spent)</h2><div class="chartbox"><canvas id="ch_chan"></canvas></div></div>
    <div class="card c12"><h2>Category breakdown</h2><div class="wrap" style="max-height:380px"><table id="cattable"></table></div></div>
  </div>
  <div class="card">
    <h2>Transactions</h2>
    <div class="filters" style="margin-bottom:10px">
      <input id="q" placeholder="Search description / party..." size="30">
      <select id="fcat"></select>
      <select id="ftype"><option value="">Debit + Credit</option><option>Debit</option><option>Credit</option></select>
      <span class="sub" id="count"></span>
    </div>
    <div class="wrap"><table id="tx"></table></div>
    <p><button id="more">Show more</button></p>
  </div>
</main>
<script id="bankdata" type="application/json">__DATA__</script>
<script>
const D = JSON.parse(document.getElementById("bankdata").textContent);
const PAL=["#2563eb","#f59e0b","#10b981","#ef4444","#8b5cf6","#06b6d4","#ec4899","#84cc16","#f97316","#14b8a6","#6366f1","#a3a3a3"];
const inr=n=>"\u20B9"+Math.round(n).toLocaleString("en-IN");
const short=n=>{const a=Math.abs(n);return (n<0?"-":"")+"\u20B9"+(a>=1e7?(a/1e7).toFixed(1)+"Cr":a>=1e5?(a/1e5).toFixed(1)+"L":a>=1e3?(a/1e3).toFixed(0)+"k":a.toFixed(0))};
const months=[...new Set(D.map(r=>r.m))].sort();
const $=id=>document.getElementById(id);
for(const id of ["from","to"]) $(id).innerHTML=months.map(m=>`<option>${m}</option>`).join("");
$("from").value=months[0]; $("to").value=months[months.length-1];
const cats=[...new Set(D.map(r=>r.c))].sort();
$("fcat").innerHTML='<option value="">All categories</option>'+cats.map(c=>`<option>${c}</option>`).join("");
const charts={};
function mk(id,cfg){if(typeof Chart==="undefined"){const b=$("err");b.style.display="block";b.textContent="Charts library could not load (tables still work).";return}if(charts[id])charts[id].destroy();charts[id]=new Chart($(id),cfg)}
function sumBy(rows,key,val){const o={};rows.forEach(r=>{o[r[key]]=(o[r[key]]||0)+r[val]});return o}
function topN(o,n){return Object.entries(o).sort((a,b)=>b[1]-a[1]).slice(0,n)}
let shown=300,filtered=[];

function current(){
  const f=$("from").value,t=$("to").value,inc=$("incself").checked;
  return D.filter(r=>r.m>=f&&r.m<=t&&(inc||r.k));
}
function render(){
  const rows=current(), mlist=months.filter(m=>m>=$("from").value&&m<=$("to").value);
  const rec=rows.reduce((a,r)=>a+r.cr,0), sp=rows.reduce((a,r)=>a+r.dr,0), net=rec-sp, nm=Math.max(mlist.length,1);
  const maxExp=rows.reduce((a,r)=>Math.max(a,r.dr),0);
  $("period").textContent=`${mlist[0]} \u2192 ${mlist[mlist.length-1]}  \u2022  ${rows.length} transactions`;
  $("kpis").innerHTML=[
    ["Total received",inr(rec),"in"],["Total spent",inr(sp),"out"],["Net saved",inr(net),net>=0?"in":"out"],
    ["Savings rate",rec?(net/rec*100).toFixed(1)+"%":"-",""],["Avg monthly spend",inr(sp/nm),""],
    ["Avg monthly income",inr(rec/nm),""],["Largest expense",inr(maxExp),""]
  ].map(k=>`<div class="card kpi"><div class="l">${k[0]}</div><div class="v ${k[2]}">${k[1]}</div></div>`).join("");

  // monthly
  const mi=sumBy(rows,"m","cr"), mo=sumBy(rows,"m","dr");
  mk("ch_month",{data:{labels:mlist,datasets:[
    {type:"bar",label:"Received",data:mlist.map(m=>mi[m]||0),backgroundColor:"#16a34a"},
    {type:"bar",label:"Spent",data:mlist.map(m=>mo[m]||0),backgroundColor:"#dc2626"},
    {type:"line",label:"Net",data:mlist.map(m=>(mi[m]||0)-(mo[m]||0)),borderColor:"#2563eb",tension:.25}]},
    options:{maintainAspectRatio:false,scales:{y:{ticks:{callback:short}}},plugins:{tooltip:{callbacks:{label:c=>c.dataset.label+": "+inr(c.parsed.y)}}}}});

  // category doughnut
  const spRows=rows.filter(r=>r.dr>0), cs=topN(sumBy(spRows,"c","dr"),99);
  const topC=cs.slice(0,10), rest=cs.slice(10).reduce((a,b)=>a+b[1],0);
  if(rest>0) topC.push(["Other",rest]);
  mk("ch_cat",{type:"doughnut",data:{labels:topC.map(x=>x[0]),datasets:[{data:topC.map(x=>x[1]),backgroundColor:PAL}]},
    options:{maintainAspectRatio:false,plugins:{legend:{position:"right",labels:{boxWidth:10,font:{size:11}}},
    tooltip:{callbacks:{label:c=>c.label+": "+inr(c.parsed)}}}}});

  // category x month stacked
  const t8=cs.slice(0,8).map(x=>x[0]);
  const ds=t8.map((c,i)=>({label:c,backgroundColor:PAL[i],data:mlist.map(m=>spRows.filter(r=>r.m===m&&r.c===c).reduce((a,r)=>a+r.dr,0))}));
  ds.push({label:"Other",backgroundColor:"#a3a3a3",data:mlist.map(m=>spRows.filter(r=>r.m===m&&!t8.includes(r.c)).reduce((a,r)=>a+r.dr,0))});
  mk("ch_catmonth",{type:"bar",data:{labels:mlist,datasets:ds},options:{maintainAspectRatio:false,
    scales:{x:{stacked:true},y:{stacked:true,ticks:{callback:short}}},plugins:{tooltip:{callbacks:{label:c=>c.dataset.label+": "+inr(c.parsed.y)}}}}});

  // payees / payers
  const hb=(id,list,color)=>mk(id,{type:"bar",data:{labels:list.map(x=>x[0]),datasets:[{data:list.map(x=>x[1]),backgroundColor:color}]},
    options:{indexAxis:"y",maintainAspectRatio:false,plugins:{legend:{display:false},tooltip:{callbacks:{label:c=>inr(c.parsed.x)}}},scales:{x:{ticks:{callback:short}}}}});
  hb("ch_payees",topN(sumBy(spRows,"p","dr"),15),"#dc2626");
  const crRows=rows.filter(r=>r.cr>0);
  hb("ch_payers",topN(sumBy(crRows,"p","cr"),15),"#16a34a");

  const ic=topN(sumBy(crRows,"c","cr"),99);
  mk("ch_inc",{type:"doughnut",data:{labels:ic.map(x=>x[0]),datasets:[{data:ic.map(x=>x[1]),backgroundColor:PAL}]},
    options:{maintainAspectRatio:false,plugins:{legend:{position:"right",labels:{boxWidth:10,font:{size:11}}},tooltip:{callbacks:{label:c=>c.label+": "+inr(c.parsed)}}}}});
  const ch=topN(sumBy(spRows,"ch","dr"),99);
  mk("ch_chan",{type:"bar",data:{labels:ch.map(x=>x[0]),datasets:[{data:ch.map(x=>x[1]),backgroundColor:"#2563eb"}]},
    options:{maintainAspectRatio:false,plugins:{legend:{display:false},tooltip:{callbacks:{label:c=>inr(c.parsed.y)}}},scales:{y:{ticks:{callback:short}}}}});

  // category table
  const tot=sp||1, cnt={};spRows.forEach(r=>cnt[r.c]=(cnt[r.c]||0)+1);
  $("cattable").innerHTML="<tr><th>Category</th><th class=n>Spent</th><th class=n>% of spend</th><th class=n>Txns</th></tr>"+
    cs.map(([c,v])=>`<tr><td>${c}</td><td class=n>${inr(v)}</td><td class=n>${(v/tot*100).toFixed(1)}%</td><td class=n>${cnt[c]}</td></tr>`).join("");
  renderTx();
}
function renderTx(){
  const q=$("q").value.toLowerCase(),c=$("fcat").value,t=$("ftype").value;
  filtered=current().filter(r=>(!c||r.c===c)&&(!t||(t==="Debit"?r.dr>0:r.cr>0))&&(!q||(r.d+" "+r.p).toLowerCase().includes(q)));
  $("count").textContent=`${filtered.length} rows \u2022 out ${inr(filtered.reduce((a,r)=>a+r.dr,0))} \u2022 in ${inr(filtered.reduce((a,r)=>a+r.cr,0))}`;
  const esc=s=>String(s).replace(/[&<>]/g,m=>({"&":"&amp;","<":"&lt;",">":"&gt;"}[m]));
  $("tx").innerHTML="<tr><th>Date</th><th>Description</th><th>Party</th><th>Category</th><th>Channel</th><th class=n>Debit</th><th class=n>Credit</th></tr>"+
    filtered.slice(0,shown).map(r=>`<tr><td>${r.dt}</td><td>${esc(r.d)}</td><td>${esc(r.p)}</td><td><span class=pill>${esc(r.c)}</span></td><td>${r.ch}</td>`+
    `<td class="n out">${r.dr?inr(r.dr):""}</td><td class="n in">${r.cr?inr(r.cr):""}</td></tr>`).join("");
  $("more").style.display=filtered.length>shown?"":"none";
}
["from","to","incself"].forEach(i=>$(i).addEventListener("change",()=>{shown=300;render()}));
["q","fcat","ftype"].forEach(i=>$(i).addEventListener("input",()=>{shown=300;renderTx()}));
$("more").onclick=()=>{shown+=300;renderTx()};
render();
</script></body></html>"""


CHARTJS_URLS = [
    "https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js",
    "https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js",
]


def get_chartjs_tag() -> str:
    """Download Chart.js once and inline it so the dashboard works offline.
    Falls back to a normal <script src> tag if the download fails."""
    import urllib.request
    for url in CHARTJS_URLS:
        try:
            with urllib.request.urlopen(url, timeout=20) as r:
                code = r.read().decode("utf-8")
            if "Chart" in code and len(code) > 50000:
                return "<script>" + code.replace("</script", "<\\/script") + "</script>"
        except Exception:
            continue
    print("  ! could not download Chart.js to embed it; dashboard will load it from the internet instead")
    return f'<script src="{CHARTJS_URLS[0]}"></script>'


def write_dashboard(df: pd.DataFrame, path: Path):
    recs = [{
        "dt": r.date.strftime("%d-%b-%Y"), "m": r.month, "d": r.desc, "p": r.party, "c": r.category, "ch": r.channel,
        "dr": round(float(r.debit), 2), "cr": round(float(r.credit), 2), "k": bool(r.counted),
    } for r in df.sort_values("date", ascending=False).itertuples()]
    # ensure_ascii + escaped < > & => the JSON can never break the HTML/JS around it
    blob = (json.dumps(recs, ensure_ascii=True, allow_nan=False)
            .replace("<", "\\u003c").replace(">", "\\u003e").replace("&", "\\u0026"))
    html = HTML.replace("__CHARTJS__", get_chartjs_tag()).replace("__DATA__", blob)
    path.write_text(html, encoding="utf-8")


# =============================================================================
# MAIN
# =============================================================================
def main():
    ap = argparse.ArgumentParser(description="Categorize a bank statement and build a report + dashboard.")
    ap.add_argument("file", help="statement .xlsx (password protected is fine)")
    ap.add_argument("-p", "--password", help="file password (or env BANK_XLSX_PASSWORD, or you'll be prompted)")
    ap.add_argument("-o", "--outdir", default="bank_report", help="output folder (default: bank_report)")
    ap.add_argument("--sheet", help="read only this sheet name")
    args = ap.parse_args()

    path = Path(args.file)
    if not path.exists():
        sys.exit(f"File not found: {path}")
    pwd = args.password or os.environ.get("BANK_XLSX_PASSWORD") or getpass.getpass("Excel password (Enter if none): ")

    print("Opening file ...")
    buf = open_workbook_bytes(path, pwd)
    print("Reading transactions ...")
    df = load_statement(buf, args.sheet)
    print("Categorizing ...")
    df = enrich(df)
    tables = build_tables(df)

    out = Path(args.outdir)
    out.mkdir(parents=True, exist_ok=True)
    write_excel(df, tables, out / "bank_report.xlsx")
    write_dashboard(df, out / "dashboard.html")
    df.drop(columns=["counted"]).to_csv(out / "transactions_clean.csv", index=False)

    c = df[df["counted"]]
    print("\n==================== DONE ====================")
    print(f"Period     : {df['date'].min():%d %b %Y} -> {df['date'].max():%d %b %Y}  ({df['month'].nunique()} months, {len(df)} txns)")
    print(f"Received   : {c['credit'].sum():,.2f}")
    print(f"Spent      : {c['debit'].sum():,.2f}")
    print(f"Net        : {c['credit'].sum() - c['debit'].sum():,.2f}")
    print("\nTop spend categories:")
    print(tables["Spend by Category"].head(8)[["Category", "Total", "% of Spend"]].to_string(index=False,
          formatters={"Total": "{:,.0f}".format, "% of Spend": "{:.1%}".format}))
    print(f"\nFiles in: {out.resolve()}")
    print("  bank_report.xlsx  |  dashboard.html  |  transactions_clean.csv")
    print("TIP: open the 'Review & Refine' sheet, then add rules to CUSTOM_RULES at the top of this script and re-run.")


if __name__ == "__main__":
    main()
